"""生成虚构企业知识语料,覆盖 4 种格式:PDF / Word(docx) / Markdown / HTML。

虚构公司:晨光科技有限公司。语料 10 份:
  - employee_handbook.md        员工手册
  - travel_reimbursement.pdf    差旅报销制度
  - product_user_guide.docx     晨光云文档用户手册
  - faq.html                    常见问题 FAQ
  - security_policy.md          信息安全管理制度
  - attendance_policy.md        考勤管理办法(B.2 混淆文档)
  - reimbursement_faq.md        报销与财务问答(B.2 混淆文档)
  - overseas_travel.md          境外差旅报销细则(B.2 混淆文档)
  - performance_review.md       绩效管理制度(B.2 混淆文档)
  - it_support.md               IT 服务台支持指南(B.2 混淆文档)

用法:python scripts/generate_corpus.py
输出:data/corpus/ 下的 10 个文件。

为什么加 5 份混淆文档(评估区分度):
  原来的 5 份文档主题互不相干,「命中正确文件」即 recall=1.0,三种检索策略
  无法区分。新增文档刻意制造三类难度:
    - 交叉文档:同一主题(如 VPN、报销材料)分散在两份文档 → recall@k 需同时命中两者
    - 近似陷阱:两份文档讲同一话题但数字/口径不同(境内 100 元 vs 境外 150 元),
      只有好的检索把「权威来源」提到前面,MRR / recall@1 才有区分度
    - 单一新题:只有新文档有答案,验证检索器在更大语料上的泛化
  所有内容完全虚构,仅供学习演示。
"""

import sys
from pathlib import Path

# 让脚本无论从哪个目录运行都能 import backend
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
# Windows 控制台默认 GBK,打印 emoji/生僻字会崩;统一重配置为 UTF-8
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

from backend.core import config

CORPUS_DIR = config.CORPUS_DIR

# =====================================================================
# 内容定义(每一节:标题, 正文)
# =====================================================================

EMPLOYEE_HANDBOOK = [
    ("第一章 总则",
     "本手册适用于晨光科技有限公司全体正式员工及实习生。所有员工应熟悉并遵守本手册的各项规定,"
     "如与劳动合同另有约定,以劳动合同为准。"),
    ("第二章 工作时间",
     "公司实行每周五天工作制,工作时间为上午 9:00 至下午 18:00,午休 12:00 至 13:30。"
     "为方便通勤,允许员工弹性上下班,弹性范围为前后半小时。每日需在考勤系统打卡。"),
    ("第三章 试用期",
     "新员工试用期为 3 个月,表现优秀者可提前转正,但试用期不得少于 1 个月。"
     "试用期内如不符合录用条件,公司可解除劳动合同。"),
    ("第四章 年休假",
     "入职满 1 年享受年休假 5 天;满 3 年享受 8 天;满 5 年享受 12 天,上限 15 天。"
     "年休假原则上在当年度内休完,逾期未休且未安排调休的,视为自动放弃。"),
    ("第五章 病假与医疗",
     "员工请病假需提供三级甲等医院的诊断证明或病假条。病假期间工资按当地规定执行,"
     "医疗期累计不超过国家规定。"),
    ("第六章 加班与调休",
     "公司提倡高效工作,不鼓励长期加班。确因工作需要加班的,需提前填写加班审批单。"
     "工作日加班可申请等时调休,周末加班按劳动法支付加班费。"),
    ("第七章 离职办理",
     "员工离职需提前 30 天书面通知公司。离职前需完成工作交接、归还办公设备,"
     "并结清各类借款与报销。最后一个工作日凭《离职交接单》办理离职手续。"),
]

TRAVEL_REIMBURSEMENT = [
    ("一、适用范围",
     "本制度适用于公司全体员工因公出差产生的交通、住宿、餐饮等费用的报销。"
     "出差须提前在 OA 系统提交出差申请,经部门负责人审批后方可出行。"),
    ("二、报销所需材料",
     "报销时必须提供以下材料:(一)合规的增值税发票;(二)行程单或登机牌、车票等凭证;"
     "(三)经审批的《出差审批单》;(四)如有超标项,需另附特别审批说明。材料不全不予受理。"),
    ("三、交通费用标准",
     "高铁/动车:原则上乘坐二等座,单程超过 4 小时可申请一等座。"
     "飞机:经济舱,机票需提前 3 个工作日通过公司差旅平台预订。"
     "市内交通:地铁、公交、打车合计每日上限 100 元。"),
    ("四、住宿费用标准",
     "一线城市(北上广深)住宿标准为每晚 500 元,其他城市每晚 350 元。"
     "超出标准的部分需在出差前获得部门负责人特别审批,否则超出部分自理。"),
    ("五、餐饮补贴",
     "出差期间按日发放餐饮补贴,标准为每日 100 元,不再另行报销餐费发票。"
     "出差不足半日的按半日计发。"),
    ("六、报销流程与时限",
     "出差结束后应在 3 个工作日内提交报销申请。流程为:OA 提交 → 部门负责人审批 → "
     "财务审核 → 出纳打款。审核通过后 15 个工作日内到账。"),
]

PRODUCT_USER_GUIDE = [
    ("产品简介",
     "晨光云文档是公司自研的企业级在线文档协作平台,支持在线编辑、多人实时协作、"
     "版本历史与精细化权限管理,是团队知识沉淀的统一入口。"),
    ("在线编辑",
     "新建文档支持富文本、Markdown、表格、图片与代码块。文档自动实时保存,"
     "编辑器支持离线模式,断网时自动保存到本地缓存,恢复联网后自动同步。"),
    ("多人协作",
     "可邀请成员以「可查看」「可评论」「可编辑」三种权限协作同一文档。"
     "多人同时编辑时,系统会以颜色区分不同成员的编辑内容,并支持评论与 @ 提醒。"),
    ("版本历史",
     "文档每 10 分钟自动生成一个版本快照,最多保留 100 个版本。"
     "可从版本历史中查看任意历史版本,支持一键回滚,误删内容可随时找回。"),
    ("导出与分享",
     "文档支持导出为 PDF、Word 与 Markdown 格式。分享时可为链接设置有效期、"
     "访问密码与访问人数上限,外部链接默认只读。"),
    ("常见故障处理",
     "如果文档无法保存,请先检查网络连接,确认右下角同步状态图标为绿色;"
     "若提示权限不足,请联系文档所有者或系统管理员调整权限;"
     "若页面白屏,请强制刷新或清除浏览器缓存后重试。"),
]

FAQ = [
    ("如何重置账号密码?",
     "在公司门户登录页点击「忘记密码」,输入绑定的企业邮箱或手机号,"
     "系统将发送重置链接。链接 30 分钟内有效,重置后需重新登录。"),
    ("如何开通企业邮箱?",
     "向部门主管提出申请,由行政部统一开通。开通后使用工号作为账号前缀,"
     "初始密码为身份证后六位,首次登录必须修改。"),
    ("如何申请 VPN 远程办公?",
     "在 IT 服务台提交 VPN 申请,注明用途与预计使用期限。"
     "审批通过后 IT 将发放一次性激活码,一个激活码仅限一人使用,禁止共享。"),
    ("如何预订会议室?",
     "通过 OA 系统的会议室模块预订,支持按容量与设备筛选。"
     "预订后如需取消,请至少提前 30 分钟释放资源,方便他人使用。"),
    ("出差报销最长可以多久?",
     "出差结束后应在 3 个工作日内提交报销申请,审核通过后 15 个工作日内到账。"
     "逾期超过 30 天提交的报销申请需部门负责人书面说明原因。"),
    ("如何申请加班审批?",
     "填写加班审批单,注明加班日期、时长与原因,由部门负责人审批。"
     "工作日加班可申请等时调休,周末加班按劳动法支付加班费。"),
]

SECURITY_POLICY = [
    ("一、口令管理",
     "员工账号口令必须不少于 12 位,并同时包含大小写字母、数字与特殊字符。"
     "口令每 90 天必须更换一次,不得与历史 5 次口令重复,禁止在多个系统间复用同一口令。"),
    ("二、数据分级",
     "公司数据按敏感程度分为三级:公开(对外宣传资料)、内部(员工通讯录、内部制度)、"
     "机密(客户数据、财务数据、源代码、商业计划)。机密数据禁止外发、禁止存储于个人设备。"),
    ("三、机密数据处理",
     "处理机密数据必须使用公司发放的加密笔记本或经 IT 批准的沙箱环境。"
     "打印机密文件需登记,废弃的机密文件必须用碎纸机销毁。"),
    ("四、终端与网络",
     "公司电脑必须安装统一的安全软件并保持自动更新。禁止在未经批准的设备上处理机密数据,"
     "禁止使用个人 VPN 绕过公司网络出口。"),
    ("五、安全事件上报",
     "发现疑似钓鱼邮件、异常登录、数据泄露等安全事件,必须在 1 小时内上报 IT 安全团队,"
     "并保留现场证据,不得自行删除相关日志。"),
]

# ── B.2 新增 5 份混淆文档 ────────────────────────────────────────

ATTENDANCE_POLICY = [
    ("一、考勤时间",
     "公司实行标准工时制,工作日上班时间为上午 8:30,下班时间为下午 17:30,"
     "午休 12:00 至 13:00。员工需每日在考勤机打卡两次:上班打卡与下班打卡。"),
    ("二、弹性工时",
     "为缓解通勤压力,考勤系统允许前后 15 分钟弹性打卡:最晚上班打卡为 8:45,"
     "最早下班打卡为 17:15。弹性范围以考勤系统设置为准。"),
    ("三、迟到早退与旷工",
     "上班打卡晚于 8:45 记为迟到,迟到 30 分钟以上按旷工半天处理;早退与迟到同规则。"
     "当月累计迟到 3 次以上,考勤评分扣减。"),
    ("四、请假审批",
     "事假、病假、年假等均须提前在 OA 系统提交请假申请,经部门负责人审批后方可休假。"
     "紧急情况可事后 24 小时内补办手续,否则按旷工处理。"),
    ("五、加班考勤",
     "加班须先填写加班审批单,并在考勤系统登记加班打卡记录。"
     "加班费以考勤打卡记录为准核发,无打卡记录的不予计算。"),
]

REIMBURSEMENT_FAQ = [
    ("报销需要什么发票?",
     "合规发票包括增值税专用发票与普通发票;电子发票须打印为 A4 纸并附国家税务总局"
     "查验水印。个人抬头的打车票、餐费发票不可报销。"),
    ("发票丢了还能报销吗?",
     "纸质发票遗失原则上不予报销;电子发票可在发票平台重新下载打印。"
     "确需报销的,需提供消费记录截图并填写《发票遗失说明》,经财务总监审批。"),
    ("报销单怎么填写?",
     "报销单须填写费用类型、发生日期、金额(含税)与发票号码,并附审批通过的"
     "《出差审批单》。费用类型选择错误会被退回。"),
    ("餐饮补贴怎么发放?",
     "出差餐饮补贴随当月工资一同发放,不再单独报销餐费。"
     "补贴标准以《差旅报销制度》为准,发放明细可在工资条中查看。"),
    ("报销为什么被退回?",
     "常见退回原因:发票抬头错误(须为公司全称)、金额与审批单不符、"
     "缺少审批单附件、费用类型填错。被退回后修改并重新提交即可。"),
    ("报销款多久到账?",
     "财务审核通过后,报销款一般在 10 个工作日内到账;涉及跨行转账最长 15 个工作日,"
     "可在财务系统查看打款进度。"),
]

OVERSEAS_TRAVEL = [
    ("一、适用范围与审批",
     "本细则适用于因公赴港澳台及境外国家的出差。境外出差须提前 15 个工作日提交申请,"
     "除部门负责人审批外,还需总经理审批,并购买境外商务旅行保险。"),
    ("二、住宿标准",
     "境外出差住宿按城市类别执行:一线国际城市(纽约、伦敦、东京、新加坡等)每晚 800 元,"
     "其他境外城市每晚 600 元。超出部分自理,特殊情况须提前特别审批。"),
    ("三、餐饮与补贴",
     "境外出差餐饮补贴为每日 150 元,无需提供餐票,随工资发放。"
     "境外出差期间不再按境内标准(每日 100 元)计发。"),
    ("四、交通与机票",
     "国际机票须提前 7 个工作日通过公司差旅平台预订,经济舱。"
     "境外市内交通凭票据实报销,每日上限 200 元。"),
    ("五、护照签证与保险",
     "员工须确保护照有效期在 6 个月以上,签证由行政部统一代办。"
     "出行前必须购买境外旅行保险,保险单据随报销材料提交。"),
]

PERFORMANCE_REVIEW = [
    ("一、考核周期",
     "公司实行半年度绩效考核,每年 1 月与 7 月各开展一次,考核范围为上一考核周期的"
     "工作表现。年度总评由两次半年度考核综合得出。"),
    ("二、评分等级",
     "绩效等级分为五档:S(卓越,占比 10%)、A(优秀,占比 20%)、B(良好,占比 50%)、"
     "C(待改进,占比 15%)、D(不合格,占比 5%)。各等级比例由部门负责人按正态分布控制。"),
    ("三、绩效与薪酬",
     "绩效等级与年度调薪和年终奖金直接挂钩:S/A 级可获上浮调薪与高额奖金,"
     "连续两次 D 级可能触发改进计划或调岗。"),
    ("四、试用期转正评估",
     "试用期员工在试用期满前一周参加转正评估,由部门负责人与导师共同评价,"
     "通过后方可转为正式员工。评估不通过的,延长试用期或解除劳动合同。"),
    ("五、申诉机制",
     "员工对绩效结果有异议的,可在考核结果公布后 5 个工作日内向人力资源部提出书面申诉,"
     "人力资源部将在 10 个工作日内组织复核。"),
]

IT_SUPPORT = [
    ("如何申请 VPN?",
     "在 IT 服务台(OA 的「IT 帮助台」入口)提交 VPN 申请,填写用途、时长与设备信息。"
     "经部门负责人与 IT 双重审批(约 2 个工作日)后发放一次性激活码,激活码仅限本人使用,禁止共享。"),
    ("如何开通企业邮箱?",
     "通过 IT 服务台提交企业邮箱开通申请,工号即邮箱账号前缀。"
     "开通后首次登录须修改初始密码,并绑定手机号开启两步验证。"),
    ("忘记账号密码怎么办?",
     "可在登录页使用「找回密码」自助重置,或联系 IT 服务台重置。"
     "重置后发送临时密码,临时密码 24 小时内有效,首次登录强制修改。"),
    ("电脑如何申请软件安装?",
     "在 IT 服务台提交软件安装申请,注明软件名称与用途。"
     "IT 审核软件版权与安全风险后远程安装,个人不得私自安装未经审批的软件。"),
    ("常见故障排查",
     "网络不通:先检查网线/WiFi,再尝试重启路由器;系统卡顿:重启并清理临时文件;"
     "蓝屏:记录错误代码后提交 IT 工单。故障处理按工单编号先后顺序。"),
    ("办公设备领用",
     "新电脑、显示器等设备通过 IT 服务台申请领用,按岗位标准配置。"
     "离职时需归还全部设备,设备损坏照价赔偿。"),
]

# =====================================================================
# 表驱动语料清单:加文档只需在 CORPUS 增一行
# (文件名, 格式, 标题, 章节列表)
# =====================================================================

CORPUS = [
    ("employee_handbook.md",     "md",   "晨光科技员工手册",           EMPLOYEE_HANDBOOK),
    ("travel_reimbursement.pdf", "pdf",  "晨光科技差旅报销制度(试行)", TRAVEL_REIMBURSEMENT),
    ("product_user_guide.docx",  "docx", "晨光云文档用户手册",         PRODUCT_USER_GUIDE),
    ("faq.html",                 "html", "晨光科技常见问题 FAQ",       FAQ),
    ("security_policy.md",       "md",   "晨光科技信息安全管理制度",   SECURITY_POLICY),
    # B.2 混淆文档:制造交叉/近似陷阱,让三种检索策略可区分
    ("attendance_policy.md",     "md",   "晨光科技考勤管理办法",       ATTENDANCE_POLICY),
    ("reimbursement_faq.md",     "md",   "晨光科技报销与财务问答",     REIMBURSEMENT_FAQ),
    ("overseas_travel.md",       "md",   "晨光科技境外差旅报销细则",   OVERSEAS_TRAVEL),
    ("performance_review.md",    "md",   "晨光科技绩效管理制度",       PERFORMANCE_REVIEW),
    ("it_support.md",            "md",   "晨光科技 IT 服务台支持指南", IT_SUPPORT),
]

# =====================================================================
# 各格式生成
# =====================================================================


def _md(filename: str, title: str, sections: list[tuple[str, str]]) -> Path:
    lines = [f"# {title}", ""]
    for heading, text in sections:
        lines += [f"## {heading}", "", text, ""]
    path = CORPUS_DIR / filename
    path.write_text("\n".join(lines), encoding="utf-8")
    print(f"✅ {filename} ({len(sections)} 节)")
    return path


def _pdf(filename: str, title: str, sections: list[tuple[str, str]]) -> Path:
    """用 reportlab 内置 STSong-Light CID 字体渲染中文 PDF,跨平台无需外部字体。

    invariant=1 必须保留:reportlab 默认把「当前时间」写进 PDF 的 CreationDate,
    同一内容每次生成字节都不同 → 内容寻址(sha256)失效 → 重复建库会污染 manifest。
    invariant=1 会固定时间戳,保证「内容不变 → doc_id 不变」的幂等语义。
    """
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    base = getSampleStyleSheet()
    title_style = ParagraphStyle("CnTitle", parent=base["Title"], fontName="STSong-Light",
                                 fontSize=18, leading=24)
    heading_style = ParagraphStyle("CnHeading", parent=base["Heading2"], fontName="STSong-Light",
                                   fontSize=13, leading=20)
    body_style = ParagraphStyle("CnBody", parent=base["BodyText"], fontName="STSong-Light",
                                fontSize=10.5, leading=18)

    story = [Paragraph(title, title_style), Spacer(1, 16)]
    for heading, text in sections:
        story.append(Paragraph(heading, heading_style))
        story.append(Paragraph(text, body_style))
        story.append(Spacer(1, 10))

    path = CORPUS_DIR / filename
    SimpleDocTemplate(str(path), pagesize=A4, title=title, invariant=1).build(story)
    print(f"✅ {filename} ({len(sections)} 节)")
    return path


def _docx(filename: str, title: str, sections: list[tuple[str, str]]) -> Path:
    import docx

    doc = docx.Document()
    doc.add_heading(title, 0)
    for heading, text in sections:
        doc.add_heading(heading, level=1)
        doc.add_paragraph(text)
    path = CORPUS_DIR / filename
    doc.save(path)
    print(f"✅ {filename} ({len(sections)} 节)")
    return path


def _html(filename: str, title: str, sections: list[tuple[str, str]]) -> Path:
    """手写语义化 HTML,BSHTMLLoader 可直接抽取正文。"""
    body_items = "".join(
        f"<h2>{h}</h2><p>{t}</p>" for h, t in sections
    )
    html = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head><meta charset="utf-8"><title>{title}</title></head>
<body>
<h1>{title}</h1>
{body_items}
</body>
</html>"""
    path = CORPUS_DIR / filename
    path.write_text(html, encoding="utf-8")
    print(f"✅ {filename} ({len(sections)} 节)")
    return path


GENERATORS = {"md": _md, "pdf": _pdf, "docx": _docx, "html": _html}


def main() -> None:
    CORPUS_DIR.mkdir(parents=True, exist_ok=True)
    for filename, fmt, title, sections in CORPUS:
        GENERATORS[fmt](filename, title, sections)
    print(f"\n🎉 已生成 {len(list(CORPUS_DIR.glob('*')))} 份语料 → {CORPUS_DIR}")


if __name__ == "__main__":
    main()
